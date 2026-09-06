import argparse
import os
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete

from app.catalog import COLOR_CODES, CORRUGATIONS, FIBERS
from app.database import Base, SessionLocal, engine
from app.models import AppSetting, CableVariant, ComponentPrice, ExchangeRate

DEFAULT_DATA_DIR = Path(__file__).resolve().parents[2] / "data"
DATA_DIR = Path(os.environ.get("TELCORD_DATA_DIR", DEFAULT_DATA_DIR)).resolve()
DEFAULT_PRICE_FILE = DATA_DIR / "telcord_cables_new.xlsx"
COMPONENT_PRICE_FILE = DATA_DIR / "telcord_prices_full.xlsx"
FIBER_CODES = {
    "SM G.652.D": "G652D", "SM G.655": "G655", "SM G.657.A1": "G657A1",
    "SM G.657.A2": "G657A2", "SM G.657.B3": "G657B3",
    "OM1": "OM1", "OM2": "OM2", "OM3": "OM3", "OM4": "OM4", "OM5": "OM5",
}
CONNECTOR_CODES = {"MTRJ Male": "MTRJ_M", "MTRJ Female": "MTRJ_F"}

def price_or_none(value: object) -> Decimal | None:
    if value is None or (isinstance(value, str) and value.strip().lower() == "нет цены"):
        return None
    return Decimal(str(value))

def available_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in {"true", "истина", "1"}:
        return True
    if isinstance(value, str) and value.strip().lower() in {"false", "ложь", "0"}:
        return False
    raise ValueError(f"Некорректное значение available: {value!r}")

def cable_key(fiber: str, product_type: str, construction: str, jacket: str | None, color: str | None) -> str:
    family = "MM" if fiber.startswith("OM") else "SM"
    construction_key = construction.upper().replace(".", "").replace("X", "X")
    return "-".join((family, fiber, product_type, construction_key, jacket or "NA", COLOR_CODES.get(color, "NA")))

def read_cable_variants(path: Path) -> list[CableVariant]:
    if not path.exists():
        raise FileNotFoundError(f"Файл цен не найден: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Кабели" not in workbook.sheetnames:
        raise ValueError("В файле отсутствует лист «Кабели»")
    ws = workbook["Кабели"]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    required = ("key", "available", "Тип изделия", "Волокно", "Конструкция", "Оболочка", "Цвет", "Код цвета", "Цена", "Валюта", "Источник цены")
    if not all(column in headers for column in required):
        raise ValueError("Лист «Кабели» не соответствует ожидаемому формату")
    pos = {name: headers.index(name) for name in headers if name}
    variants: list[CableVariant] = []
    for row_number, row in enumerate(rows, 2):
        key = str(row[pos["key"]]).strip()
        product_type = row[pos["Тип изделия"]]
        component = row[pos["Волокно"]]
        construction = row[pos["Конструкция"]]
        jacket = row[pos["Оболочка"]]
        color = row[pos["Цвет"]]
        raw_price = price_or_none(row[pos["Цена"]])
        if raw_price is None:
            continue
        currency = row[pos["Валюта"]] if raw_price is not None else None
        source_value = row[pos["Источник цены"]]
        source = f"{path.name}; лист Кабели; строка {row_number}; {source_value or 'источник цены не указан'}"
        fiber = FIBER_CODES.get(component)
        if not fiber or fiber not in FIBERS:
            raise ValueError(f"Строка {row_number}: неизвестное волокно {component}")
        variants.append(CableVariant(
            key=key, product_type=product_type, fiber_type=fiber, fiber_name=component,
            construction=construction, jacket=jacket, color=color,
            color_code=row[pos["Код цвета"]], raw_price=raw_price, currency=currency,
            source=source, available=available_bool(row[pos["available"]]),
        ))
    keys = [variant.key for variant in variants]
    if len(keys) != len(set(keys)):
        raise ValueError("На листе «Кабели» обнаружены дублирующиеся key")
    return variants

def read_component_prices(path: Path) -> list[ComponentPrice]:
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    pos = {name: headers.index(name) for name in headers if name}
    components: list[ComponentPrice] = []
    for row_number, row in enumerate(rows, 2):
        group = row[pos["Группа"]]
        component = row[pos["Волокно / компонент"]]
        construction = row[pos["Конструкция"]]
        raw_price = price_or_none(row[pos["Цена"]])
        currency = row[pos["Валюта"]]
        source = f"{path.name}; лист Цены; строка {row_number}; {row[pos['Источник']] or 'источник не указан'}"
        if group == "Коннектор":
            components.append(ComponentPrice(component_type="connector", component_code=f"{component}:{construction}", description=f"{component}/{construction}", raw_price=raw_price, currency=currency, source=source))
        elif group == "MTRJ":
            connector = CONNECTOR_CODES.get(component)
            if not connector:
                raise ValueError(f"Строка {row_number}: неизвестный MTRJ")
            for polish in ("UPC", "PC"):
                components.append(ComponentPrice(component_type="connector", component_code=f"{connector}:{polish}", description=f"{component}/{polish}", raw_price=raw_price, currency=currency, source=source))
        elif group == "Holder":
            components.append(ComponentPrice(component_type="holder", component_code=str(component), description=f"{component} duplex holder", raw_price=raw_price, currency=currency, source=source))
        elif group == "Гофра" and construction in CORRUGATIONS:
            components.append(ComponentPrice(component_type="corrugation", component_code=str(construction), description=str(component), raw_price=raw_price, currency=currency, source=source))
    return components

def read_workbook(path: Path) -> tuple[list[CableVariant], list[ComponentPrice]]:
    return read_cable_variants(path), read_component_prices(COMPONENT_PRICE_FILE)

def seed(excel: Path = DEFAULT_PRICE_FILE) -> None:
    variants, components = read_workbook(excel)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        db.execute(delete(CableVariant)); db.execute(delete(ComponentPrice))
        db.execute(delete(ExchangeRate)); db.execute(delete(AppSetting))
        db.add_all(variants); db.add_all(components)
        for price in list(components):
            if price.component_type in {"connector", "holder"}:
                db.add(ComponentPrice(component_type=price.component_type, component_code=f"{price.component_code}:BK", description=f"{price.description} black", raw_price=None, currency=price.currency or "USD", source=f"{excel.name}; цена full_black ожидается"))
        db.add(ExchangeRate(pair="USD/RUB", rate=Decimal("90")))
        db.add_all([
            AppSetting(key="exchange_rate_usd_rub", value="90"),
            AppSetting(key="production_status", value="NORMAL"),
            AppSetting(key="production_days_min", value="2"),
            AppSetting(key="production_days_max", value="3"),
            AppSetting(key="production_title", value="Обычная загрузка производства"),
            AppSetting(key="production_message", value="Ориентировочный срок изготовления при заказе сегодня. Фактический срок зависит от количества."),
        ])
        db.commit()

if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--excel", type=Path, default=DEFAULT_PRICE_FILE)
    seed(parser.parse_args().excel)
